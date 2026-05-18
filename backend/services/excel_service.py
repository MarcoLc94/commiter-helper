import io
from datetime import date as date_type
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from models import ReportRequest

C_HEADER_BG = "1F497D"
C_HEADER_FG = "FFFFFF"
C_WEEKEND_BG = "C00000"
C_WEEKEND_FG = "FFFFFF"
C_HOLIDAY_BG = "E26B0A"
C_HOLIDAY_FG = "FFFFFF"
C_DATE_BG = "BDD7EE"
C_SUBTOTAL_BG = "DEEAF1"
C_SUBTOTAL_FG = "1F497D"
C_WHITE = "FFFFFF"

WEEKDAYS_ES = ["Lunes", "Martes", "Miercoles", "Jueves", "Viernes", "Sabado", "Domingo"]


def fill(color: str) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def font(bold=False, color="000000", size=10, italic=False) -> Font:
    return Font(name="Calibri", bold=bold, color=color, size=size, italic=italic)


def align(h="left", v="top", wrap=True) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def date_label(date_str: str, weekday: str) -> str:
    d = date_type.fromisoformat(date_str)
    return f"{weekday} {d.day:02d}/{d.month:02d}/{d.year}"


def style_row(ws, row_num: int, bg: str, fg: str, bold: bool, h_align: str, values: list, merge_from=None):
    for col, val in enumerate(values, 1):
        cell = ws.cell(row=row_num, column=col, value=val)
        cell.fill = fill(bg)
        cell.font = font(bold=bold, color=fg)
        cell.alignment = align(h=h_align, v="center", wrap=False)
    if merge_from:
        ws.merge_cells(start_row=row_num, start_column=merge_from, end_row=row_num, end_column=7)


def sanitize_sheet_name(name: str) -> str:
    invalid = r"/\?*:[]"
    for ch in invalid:
        name = name.replace(ch, "-")
    return name[:31]


def generate_excel(request: ReportRequest) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = sanitize_sheet_name(request.sheet_name)

    col_widths = [16, 32, 20, 12, 7, 40, 13]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Header
    ws.row_dimensions[1].height = 30
    headers = ["Fecha", "Nombre del Ticket / Actividad", "Departamento / Modulo",
               "Status", "Horas", "Comentarios de Cierre", "Atendio"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = fill(C_HEADER_BG)
        cell.font = font(bold=True, color=C_HEADER_FG, size=11)
        cell.alignment = align(h="center", v="center")

    for day in request.days:
        label = date_label(day.date, day.weekday)
        r = ws.max_row + 1

        if day.day_type == "weekend":
            ws.row_dimensions[r].height = 22
            style_row(ws, r, C_WEEKEND_BG, C_WEEKEND_FG, True, "center",
                      [label, "FIN DE SEMANA", "", "", "", "", ""], merge_from=2)

        elif day.day_type == "holiday":
            ws.row_dimensions[r].height = 22
            holiday_label = day.activities[0].activity.upper() if day.activities else "DIA FESTIVO"
            style_row(ws, r, C_HOLIDAY_BG, C_HOLIDAY_FG, True, "center",
                      [label, holiday_label, "", "", "", "", ""], merge_from=2)

        elif day.day_type == "workday":
            if not day.activities:
                ws.row_dimensions[r].height = 22
                ws.cell(row=r, column=1, value=label).fill = fill(C_DATE_BG)
                ws.cell(row=r, column=1).font = font(bold=True)
                ws.cell(row=r, column=1).alignment = align(h="center", v="center")
                for col in range(2, 8):
                    c = ws.cell(row=r, column=col, value="")
                    c.fill = fill(C_WHITE)
                    c.font = font()
            else:
                for i, activity in enumerate(day.activities):
                    row_num = ws.max_row + 1
                    ws.row_dimensions[row_num].height = 45

                    date_val = label if i == 0 else ""
                    date_bg = C_DATE_BG if i == 0 else C_WHITE
                    date_bold = i == 0

                    dc = ws.cell(row=row_num, column=1, value=date_val)
                    dc.fill = fill(date_bg)
                    dc.font = font(bold=date_bold)
                    dc.alignment = align(h="center", v="center")

                    act_name = f"{activity.ticket} - {activity.activity}" if activity.ticket else activity.activity

                    cells = [
                        (2, act_name, False, "left"),
                        (3, activity.module, False, "left"),
                        (4, activity.status, False, "left"),
                        (5, activity.hours, True, "center"),
                        (6, activity.comments, False, "left"),
                        (7, activity.author, False, "left"),
                    ]
                    for col, val, bold, h in cells:
                        c = ws.cell(row=row_num, column=col, value=val)
                        c.fill = fill(C_WHITE)
                        c.font = font(bold=bold)
                        c.alignment = align(h=h, v="top", wrap=True)

                sub_r = ws.max_row + 1
                ws.row_dimensions[sub_r].height = 16
                total = sum(a.hours for a in day.activities)
                sc = ws.cell(row=sub_r, column=1, value=f"TOTAL DIA: {total:.1f} h")
                sc.fill = fill(C_SUBTOTAL_BG)
                sc.font = font(bold=True, color=C_SUBTOTAL_FG)
                sc.alignment = align(h="right", v="center", wrap=False)
                ws.merge_cells(start_row=sub_r, start_column=1, end_row=sub_r, end_column=7)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
